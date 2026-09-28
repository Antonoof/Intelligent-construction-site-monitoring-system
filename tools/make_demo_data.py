"""Генерация демо-данных: два проекта, календарные графики, камеры и зоны, синтетические снимки.

Снимки синтетические (ТЗ, раздел 6.2, допускает обезличенные или синтетические снимки): сцены
в стиле демо «План-факт» с разной освещённостью, погодой и сезоном. Разметка рядом со снимком
(<имя>.json) — эталонные рамки сцены; её читает демо-детектор, пока не подключены веса RF-DETR.

Каждый снимок — сценарий методики: пример из ТЗ (экскаватор без самосвалов), нехватка техники,
лишняя техника, опережение и затягивание этапов, пустая зона, «нет данных» ночью и в тумане,
расхождение стадии по кадру общего плана. Ожидаемые выводы — в data/demo/README.md и в тестах.

Запуск из корня репозитория:  python tools/make_demo_data.py
"""
from __future__ import annotations

import asyncio
import datetime as dt
import hashlib
import io
import json
import random
import sys
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageFilter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from demo_scenes import compose  # noqa: E402

OUT = ROOT / "data" / "demo"
SCALE = 4  # 320×180 → 1280×720


def norm_poly(pts):
    return [[round(x / 320, 4), round(y / 180, 4)] for x, y in pts]


# ============================================================ проект 1: жилой дом, лето

P1 = {
    "slug": "housing",
    "name": "Жилой дом, корпус 1 (демо, синтетические снимки)",
    "object_type": "housing",
    "address": "Москва, демонстрационный объект",
    "description": "Подготовка территории и подземная часть. 6 камер, 6 зон, 2 дня съёмки: 24–25.09.2026.",
    "schedule": "schedule_housing.xlsx",
    "site": {
        "zones": [
            {"key": "Z1", "name": "Котлован, секция 1", "color": "#2A78D6"},
            {"key": "Z2", "name": "Секция 2 · фундаментная плита", "color": "#EB6834"},
            {"key": "Z3", "name": "Подъездная дорога", "color": "#EDA100"},
            {"key": "Z4", "name": "Ливнёвка", "color": "#E87BA4"},
            {"key": "Z5", "name": "Стройгородок и въезд", "color": "#1BAF7A"},
            {"key": "Z6", "name": "Секция 3", "color": "#7A5CC2"},
        ],
        "cameras": [
            {"key": "CAM-01", "name": "Котлован · север", "scene": "pit", "url": "rtsp://10.20.0.11:554/stream1",
             "zones": {"Z1": norm_poly([(16, 96), (304, 96), (318, 179), (2, 179)])}},
            {"key": "CAM-02", "name": "Секция 2 · плита", "scene": "slab", "url": "http://10.20.0.12/snapshot.jpg",
             "zones": {"Z2": norm_poly([(26, 94), (296, 94), (318, 179), (2, 179)])}},
            {"key": "CAM-03", "name": "Въезд · дорога · секция 3", "scene": "entrance", "url": "rtsp://10.20.0.13:554/stream1",
             "zones": {"Z3": norm_poly([(2, 112), (146, 106), (166, 179), (2, 179)]),
                       "Z6": norm_poly([(184, 104), (318, 100), (318, 179), (192, 179)])}},
            {"key": "CAM-04", "name": "Ливнёвка", "scene": "drain", "url": "http://10.20.0.14/snapshot.jpg",
             "zones": {"Z4": norm_poly([(14, 98), (306, 98), (318, 179), (2, 179)])}},
            {"key": "CAM-05", "name": "Стройгородок", "scene": "camp", "url": "rtsp://10.20.0.15:554/stream1",
             "zones": {"Z5": norm_poly([(2, 101), (318, 101), (318, 179), (2, 179)])}},
            {"key": "CAM-06", "name": "Общий план · мачта", "scene": "overview", "overview": True, "tiles": 3,
             "url": "rtsp://10.20.0.16:554/stream1",
             "note": "Камера общего плана с высоты: детекция по фрагментам 3×3 и оценка стадии/готовности (ML-часть)",
             "zones": {"Z1": norm_poly([(20, 70), (150, 62), (170, 150), (6, 168)]),
                       "Z2": norm_poly([(184, 62), (300, 58), (316, 138), (196, 148)])}},
        ],
        "site_equipment": [],
    },
}

S1_ROWS = [  # код, наименование, зона, начало, окончание, техника, подрядчик
    ("10", "Подготовка территории", "", "01.07.2026", "15.10.2026", "", "ООО «СтройПодготовка»"),
    ("10.11", "Обустройство строительной площадки", "", "01.07.2026", "19.09.2026", "", "ООО «СтройПодготовка»"),
    ("10.11.1", "Устройство ограждения строительной площадки", "Стройгородок и въезд", "01.07.2026", "15.07.2026",
     "Кран-манипулятор ×1", "ООО «СтройПодготовка»"),
    ("10.11.4", "Возведение хозяйственно-бытового городка", "Стройгородок и въезд", "16.07.2026", "10.08.2026",
     "Кран-манипулятор ×1", "ООО «СтройПодготовка»"),
    ("", "Устройство временных дорог", "Подъездная дорога", "25.08.2026", "19.09.2026",
     "Грейдер ×1; Каток ×1; Самосвал ×2", "ООО «ДорСтрой»"),
    ("10.2", "Вынос инженерных систем из пятна застройки", "", "01.09.2026", "15.10.2026", "", "ООО «ИнжСети»"),
    ("10.2.6", "Вынос сетей: ливневая канализация", "Ливнёвка", "15.09.2026", "15.10.2026",
     "Экскаватор ×1; Самосвал ×1", "ООО «ИнжСети»"),
    ("12", "Выполнение строительно-монтажных работ", "", "07.09.2026", "30.12.2026", "", "ООО «ГенСтрой»"),
    ("12.3", "Устройство подземной части", "", "07.09.2026", "30.10.2026", "", "ООО «ГенСтрой»"),
    ("12.3.1", "Устройство котлована", "Котлован, секция 1", "07.09.2026", "09.10.2026",
     "Экскаватор ×1; Самосвал ×2", "ООО «ЗемРесурс»"),
    ("12.3.4", "Устройство фундамента", "Секция 2 · фундаментная плита", "14.09.2026", "05.10.2026", "", "ООО «МонолитСтрой»"),
    ("12.3.4.4", "Устройство бетонной подготовки", "Секция 2 · фундаментная плита", "14.09.2026", "20.09.2026",
     "Автобетоносмеситель ×1", "ООО «МонолитСтрой»"),
    ("12.3.4.6", "Устройство фундаментной плиты", "Секция 2 · фундаментная плита", "21.09.2026", "05.10.2026",
     "Автобетоносмеситель ×2; Автобетононасос ×1", "ООО «МонолитСтрой»"),
    ("12.3.4.1", "Устройство свайного фундамента", "Секция 3", "05.10.2026", "25.10.2026",
     "Буровая установка ×1", "ООО «Геофонд»"),
    ("12.3.7.2", "Обратная засыпка", "Котлован, секция 1", "20.10.2026", "30.10.2026",
     "Бульдозер ×1; Каток ×1", "ООО «ЗемРесурс»"),
    ("12.4.29", "Каркас здания", "Секция 2 · фундаментная плита", "06.10.2026", "20.12.2026",
     "Башенный кран ×1; Автобетоносмеситель ×2", "ООО «МонолитСтрой»"),
    ("", "Монтаж башенного крана", "Секция 2 · фундаментная плита", "15.09.2026", "18.09.2026", "Автокран ×1",
     "ООО «КранСервис»"),
]

# m(класс, gx, gy, масштаб, уверенность, ...): точка опоры (gx, gy) в координатах сцены 320×180
def m(cls, gx, gy, s, conf=0.9, **kw):
    return {"cls": cls, "gx": gx, "gy": gy, "s": s, "conf": conf, **kw}


S1 = [
    # ---- 24.09.2026
    ("CAM-01", "2026-09-24 08:05", "morning", "clear",
     [m("excavator", 64, 150, 1.0, 0.93, pose=0), m("dump_truck", 172, 164, .95, 0.9, flip=True, loaded=True),
      m("dump_truck", 256, 128, .72, 0.86, flip=True)],
     "Котлован: экскаватор и 2 самосвала — комплект по графику"),
    ("CAM-01", "2026-09-24 10:30", "day", "clear", [m("excavator", 118, 152, 1.05, 0.92, pose=1)],
     "Пример из ТЗ: экскаватор работает, самосвалов нет"),
    ("CAM-01", "2026-09-24 11:20", "day", "clear", [m("excavator", 150, 146, 1.0, 0.91, pose=2)],
     "Экскаватор по-прежнему без самосвалов — отклонение подтверждается серией"),
    ("CAM-01", "2026-09-24 14:00", "day", "clear",
     [m("excavator", 70, 150, 1.0, 0.92, pose=0), m("dump_truck", 196, 164, .95, 0.89, flip=True, loaded=True),
      m("roller", 262, 128, .8, 0.87)],
     "В котловане появился каток — не соответствует этапу «Разработка грунта»"),
    ("CAM-01", "2026-09-24 16:30", "dusk", "clear",
     [m("excavator", 92, 150, 1.0, 0.88, pose=1), m("dump_truck", 196, 166, .95, 0.85, flip=True),
      m("dump_truck", 262, 130, .7, 0.8, flip=True, loaded=True)],
     "Вечер: комплект восстановлен"),
    ("CAM-02", "2026-09-24 09:00", "day", "clear",
     [m("concrete_mixer", 212, 164, .95, 0.9, flip=True), m("concrete_pump", 72, 156, 1.0, label=False, up=True),
      m("tower_crane", 236, 92, .8, 0.83)],
     "Фундаментная плита: 1 миксер вместо 2 по графику, насос работает (класса «бетононасос» в модели пока нет)"),
    ("CAM-02", "2026-09-24 11:00", "day", "clear", [m("concrete_mixer", 150, 162, .95, 0.91)],
     "Миксер стоит без насоса и крана"),
    ("CAM-02", "2026-09-24 13:30", "day", "clear",
     [m("concrete_mixer", 212, 164, .95, 0.89, flip=True), m("tower_crane", 236, 92, .8, 0.81)],
     "Снова один миксер — за день не больше одного из двух"),
    ("CAM-03", "2026-09-24 09:30", "day", "rain",
     [m("grader", 62, 150, .85, 0.86), m("roller", 112, 170, .85, 0.83)],
     "Дождь. Грейдер и каток на подъездной дороге, хотя этап закончился 19.09"),
    ("CAM-03", "2026-09-24 15:10", "day", "clear",
     [m("drilling_rig", 262, 152, .72, 0.88), m("truck", 72, 166, .85, 0.9)],
     "Буровая установка в секции 3 — сваи по графику с 05.10 (опережение); грузовик на дороге — доставка"),
    ("CAM-03", "2026-09-24 16:40", "dusk", "clear", [m("mobile_crane", 250, 150, .55, 0.44)],
     "Сомнительная рамка «автокран» 0,44 — ниже порога, отклонения нет"),
    ("CAM-04", "2026-09-24 09:10", "day", "clear", [], "Ливнёвка: траншея пуста"),
    ("CAM-04", "2026-09-24 12:40", "day", "clear", [], "Ливнёвка: техники нет весь день — работы не ведутся"),
    ("CAM-04", "2026-09-24 21:40", "dark", "clear", [], "Ночь без освещения — «нет данных», а не «нет техники»"),
    ("CAM-05", "2026-09-24 10:00", "day", "clear",
     [m("truck", 112, 166, 1.0, 0.92, cargo=True), m("crane_manipulator", 236, 166, 1.0, 0.9, flip=True, up=True)],
     "Стройгородок: разгрузка материалов — доставка допустима"),
    ("CAM-05", "2026-09-24 16:00", "day", "clear", [m("excavator", 205, 152, .9, 0.87, pose=0)],
     "Экскаватор в зоне без работ по графику"),
    ("CAM-06", "2026-09-24 12:00", "day", "clear",
     [m("excavator", 82, 120, .35, 0.74, pose=1), m("dump_truck", 118, 132, .33, 0.71),
      m("concrete_mixer", 244, 118, .33, 0.69), m("tower_crane", 262, 60, .5, 0.78)],
     "Общий план: мелкая техника — детекция по фрагментам; стадия S1, готовность 11,8 %",
     {"stage": "S1", "readiness": 11.8, "model": "dinov2-vitg-head-v1 (ML-часть; в демо — из разметки)"}),
    # ---- 25.09.2026
    ("CAM-01", "2026-09-25 08:10", "morning", "clear",
     [m("excavator", 70, 150, 1.0, 0.92, pose=0), m("dump_truck", 196, 164, .95, 0.9, flip=True)],
     "Котлован: экскаватор с самосвалом"),
    ("CAM-01", "2026-09-25 09:40", "day", "clear",
     [m("excavator", 120, 150, 1.0, 0.93, pose=2), m("dump_truck", 212, 166, .95, 0.9, flip=True, loaded=True),
      m("dump_truck", 58, 124, .7, 0.84)],
     "Котлован: 2 самосвала"),
    ("CAM-02", "2026-09-25 10:00", "day", "clear",
     [m("concrete_mixer", 212, 164, .95, 0.91, flip=True), m("concrete_mixer", 150, 138, .75, 0.86, flip=True),
      m("concrete_pump", 60, 156, 1.0, label=False, up=True), m("tower_crane", 236, 92, .8, 0.82)],
     "Плита: 2 миксера, насос — по графику"),
    ("CAM-03", "2026-09-25 11:30", "day", "clear", [], "Дорога и секция 3 свободны"),
    ("CAM-04", "2026-09-25 09:00", "day", "clear",
     [m("excavator", 90, 150, 1.0, 0.9, pose=1), m("dump_truck", 236, 162, .9, 0.88, flip=True)],
     "Ливнёвка: работы начались"),
    ("CAM-04", "2026-09-25 11:00", "day", "clear", [m("excavator", 96, 150, 1.0, 0.9, pose=0)],
     "Траншея: экскаватор без самосвала — к сведению (грунт может складироваться рядом)"),
    ("CAM-05", "2026-09-25 07:20", "morning", "fog", [], "Туман — снимок не проходит контроль качества"),
    ("CAM-06", "2026-09-25 14:00", "day", "clear",
     [m("excavator", 90, 118, .35, 0.73, pose=2), m("dump_truck", 128, 134, .33, 0.7, flip=True),
      m("concrete_mixer", 236, 116, .33, 0.72), m("concrete_mixer", 262, 128, .33, 0.66, flip=True),
      m("tower_crane", 262, 60, .5, 0.8)],
     "Общий план: стадия S1, готовность 12,6 %",
     {"stage": "S1", "readiness": 12.6, "model": "dinov2-vitg-head-v1 (ML-часть; в демо — из разметки)"}),
    # добавлен в конец списка, чтобы не менять шум и разметку остальных снимков (seed зависит от номера)
    ("CAM-04", "2026-09-25 13:20", "day", "clear", [m("dump_truck", 232, 160, .9, 0.88, flip=True)],
     "Траншея: самосвал ждёт погрузки, экскаватора нет — нет необходимой техники"),
]

# ============================================================ проект 2: школа, зима

P2 = {
    "slug": "school",
    "name": "Школа на 550 мест (демо, зима)",
    "object_type": "education",
    "address": "Москва, демонстрационный объект",
    "description": "Монолитный каркас зимой: снег, ночная съёмка с прожекторами, кадр общего плана с оценкой стадии.",
    "schedule": "schedule_school.xlsx",
    "site": {
        "zones": [
            {"key": "Z1", "name": "Корпус А", "color": "#2A78D6"},
            {"key": "Z2", "name": "Площадка складирования", "color": "#EDA100"},
        ],
        "cameras": [
            {"key": "CAM-01", "name": "Общий план · мачта", "scene": "school_overview", "overview": True, "tiles": 3,
             "url": "rtsp://10.30.0.11:554/stream1",
             "zones": {"Z1": norm_poly([(100, 92), (252, 90), (266, 152), (92, 156)]),
                       "Z2": norm_poly([(8, 116), (96, 114), (92, 176), (4, 178)])}},
            {"key": "CAM-02", "name": "Корпус А · юг", "scene": "school", "url": "rtsp://10.30.0.12:554/stream1",
             "zones": {"Z1": norm_poly([(0, 124), (320, 124), (320, 179), (0, 179)])}},
        ],
        "site_equipment": [{"class": "tower_crane", "name": "Башенный кран КБ-586 №1", "zones": ["Z1", "Z2"],
                            "from": "2026-01-05"}],
    },
}

S2_ROWS = [
    ("12.3.7.2", "Обратная засыпка", "Корпус А", "15.12.2025", "25.01.2026", "Бульдозер ×1", "ООО «ЗемРесурс»"),
    ("12.4.29", "Каркас здания", "Корпус А", "12.01.2026", "30.04.2026", "Башенный кран ×1", "ООО «МонолитСтрой»"),
    ("12.4.10", "Бетонирование колонн и перекрытий", "Корпус А", "12.01.2026", "30.04.2026", "Автобетоносмеситель ×1",
     "ООО «МонолитСтрой»"),
    ("12.4.2", "Поставка оборудования", "Площадка складирования", "01.02.2026", "28.02.2026", "Грузовик ×1",
     "ООО «ТехСнаб»"),
]

S2 = [
    ("CAM-02", "2026-02-11 07:00", "dark", "snow", [], "До рассвета — «нет данных»"),
    ("CAM-02", "2026-02-11 10:15", "day", "snow",
     [m("tower_crane", 240, 132, 1.05, 0.91), m("concrete_mixer", 82, 172, 1.0, 0.9),
      m("concrete_pump", 190, 176, .9, label=False)],
     "Снег. Бетонирование: миксер и башенный кран — по графику"),
    ("CAM-01", "2026-02-11 12:00", "day", "snow",
     [m("tower_crane", 196, 98, .6, 0.84), m("concrete_mixer", 150, 132, .34, 0.71),
      m("truck", 58, 150, .38, 0.74, cargo=True)],
     "Общий план: модель готовности видит стадию S2 при плановой S3",
     {"stage": "S2", "readiness": 31.5, "model": "dinov2-vitg-head-v1 (ML-часть; в демо — из разметки)"}),
    ("CAM-02", "2026-02-11 13:40", "day", "snow",
     [m("tower_crane", 240, 132, 1.05, 0.9), m("bulldozer", 150, 174, 1.0, 0.88, up=True)],
     "Бульдозер у корпуса А: обратная засыпка по графику закончилась 25.01 (на деле — уборка снега: инженер отклонит)"),
    ("CAM-01", "2026-02-11 19:30", "night_lit", "clear", [m("tower_crane", 196, 98, .6, 0.79)],
     "Ночь, прожекторы: снимок годен, башенный кран на месте"),
]


# ============================================================ сборка

def write_schedule(path: Path, title: str, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "График"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    head = ["Код", "Наименование работ", "Захватка", "Начало", "Окончание", "Техника (план)", "Подрядчик"]
    for j, h in enumerate(head, 1):
        c = ws.cell(3, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="310F53")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, r in enumerate(rows, 4):
        code, name, zone, a, b, eq, contr = r
        vals = [code, name, zone, dt.datetime.strptime(a, "%d.%m.%Y"), dt.datetime.strptime(b, "%d.%m.%Y"), eq, contr]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            if j in (4, 5):
                c.number_format = "DD.MM.YYYY"
        level = code.count(".") if code else 1
        ws.cell(i, 2).alignment = Alignment(indent=level)
        if not zone:
            for j in range(1, 8):
                ws.cell(i, j).font = Font(bold=True)
    for col, w in zip("ABCDEFG", [10, 52, 30, 12, 12, 40, 24]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.save(path)


def add_noise(png: bytes, seed: int) -> bytes:
    img = Image.open(io.BytesIO(png)).convert("RGB").filter(ImageFilter.GaussianBlur(0.6))
    a = np.asarray(img).astype(np.float32)
    rng = np.random.default_rng(seed)
    a += rng.normal(0, 3.2, a.shape)
    out = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))
    buf = io.BytesIO()
    out.save(buf, "JPEG", quality=86)
    return buf.getvalue()


async def render(jobs):
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        b = await p.chromium.launch()
        pg = await b.new_page(viewport={"width": 1280, "height": 720})
        results = []
        for svg in jobs:
            await pg.set_content(f'<html><body style="margin:0;background:#000">{svg}</body></html>')
            await pg.wait_for_timeout(60)
            results.append(await pg.screenshot(type="png"))
        await b.close()
        return results


def build_project(p: dict, shots, rows, seed0: int):
    pdir = OUT / p["slug"]
    sdir = pdir / "snapshots"
    sdir.mkdir(parents=True, exist_ok=True)
    for old in sdir.glob("*"):
        old.unlink()
    write_schedule(pdir / p["schedule"], f'Календарный график · {p["name"]}', rows)
    cams = {c["key"]: c for c in p["site"]["cameras"]}
    svgs, metas = [], []
    for k, shot in enumerate(shots):
        cam, when, light, weather, machines, note = shot[:6]
        assessment = shot[6] if len(shot) > 6 else None
        t = dt.datetime.strptime(when, "%Y-%m-%d %H:%M")
        svg, boxes = compose(cams[cam]["scene"], light, weather, machines, cam, t.strftime("%d.%m.%Y  %H:%M:%S"),
                             cams[cam]["name"], seed=seed0 + k)
        svgs.append(svg)
        rnd = random.Random(seed0 * 100 + k)
        dets = []
        for bx in boxes:
            x1, y1, x2, y2 = bx["box"]
            w, h = x2 - x1, y2 - y1
            jit = [rnd.uniform(-.015, .015) * w, rnd.uniform(-.015, .015) * h,
                   rnd.uniform(-.015, .015) * w, rnd.uniform(-.01, .01) * h]
            xyxy = [max(0, (x1 + jit[0]) * SCALE), max(0, (y1 + jit[1]) * SCALE),
                    min(1280, (x2 + jit[2]) * SCALE), min(720, (y2 + jit[3]) * SCALE)]
            dets.append({"cls": bx["cls"], "conf": bx["conf"], "xyxy": [round(v, 1) for v in xyxy]})
        name = f'{cam}_{t:%Y-%m-%d_%H-%M}'
        metas.append((name, {"camera": cam, "taken_at": t.isoformat(), "scene": cams[cam]["scene"], "light": light,
                             "weather": weather, "note": note, "detections": dets, "assessment": assessment,
                             "source": "синтетическая сцена, рамки — эталонная разметка сцены"}))
    pngs = asyncio.run(render(svgs))
    for k, ((name, meta), png) in enumerate(zip(metas, pngs)):
        jpg = add_noise(png, seed0 + k)
        (sdir / f"{name}.jpg").write_bytes(jpg)
        meta["sha256"] = hashlib.sha256(jpg).hexdigest()
        (sdir / f"{name}.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
    site = json.loads(json.dumps(p["site"]))
    for c in site["cameras"]:
        c.pop("scene", None)
    (pdir / "site.json").write_text(json.dumps(site, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"slug": p["slug"], "name": p["name"], "object_type": p["object_type"], "address": p["address"],
            "description": p["description"], "schedule": f'{p["slug"]}/{p["schedule"]}', "site": f'{p["slug"]}/site.json',
            "snapshots": f'{p["slug"]}/snapshots'}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    projects = [build_project(P1, S1, S1_ROWS, 100), build_project(P2, S2, S2_ROWS, 900)]
    (OUT / "projects.json").write_text(json.dumps({"projects": projects}, ensure_ascii=False, indent=1),
                                       encoding="utf-8")
    n = sum(1 for _ in OUT.rglob("*.jpg"))
    print(f"демо-данные: {len(projects)} проекта, {n} снимков → {OUT}")


if __name__ == "__main__":
    main()
