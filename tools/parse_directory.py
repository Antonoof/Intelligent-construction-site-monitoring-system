"""Разбор справочника видов работ организаторов (XLSX) в плоский список с иерархией.

Особенности исходного файла, которые учитывает разбор:
* Excel превратил часть кодов («10.1», «12.3») в даты — восстанавливаем «день.месяц»;
* строки без кода — подпункты ближайшей строки с кодом, получают id «<код>-<N>»;
* отметка «˅» в столбцах C–K — применимость вида работ к типу объекта;
* столбец L — комментарии экспертов (L172:L179 объединены).
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

OBJECT_TYPES = [
    ("housing", "Жильё"), ("education", "Образование"), ("healthcare", "Здравоохранение"),
    ("sport", "Спорт"), ("culture", "Культура"), ("admin", "Административные здания"),
    ("kindergarten", "ДОУ"), ("office", "Офисно-деловой центр"), ("roads", "Дороги"),
]


def norm_code(raw) -> str | None:
    if raw is None or str(raw).strip() == "":
        return None
    if isinstance(raw, (dt.datetime, dt.date)):
        # «10.1» → 10 января: день — первая часть кода, месяц — вторая
        return f"{raw.day}.{raw.month}"
    s = str(raw).strip().rstrip(".")
    if s.endswith(".0"):
        s = s[:-2]
    return s


def parse(path: Path) -> list[dict]:
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    # комментарии с учётом объединённых ячеек столбца L
    notes: dict[int, str] = {}
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, 12).value
        if v:
            notes[r] = str(v).strip()
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 12 and rng.max_col == 12:
            v = ws.cell(rng.min_row, 12).value
            for r in range(rng.min_row, rng.max_row + 1):
                notes[r] = str(v).strip()

    rows: list[dict] = []
    by_code: dict[str, dict] = {}
    last_coded: dict | None = None
    sub_counter: dict[str, int] = {}
    for r in range(3, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if name is None or str(name).strip() == "":
            continue
        name = " ".join(str(name).split())
        code = norm_code(ws.cell(r, 1).value)
        applies = [key for i, (key, _) in enumerate(OBJECT_TYPES) if ws.cell(r, 3 + i).value not in (None, "")]
        if code:
            parts = code.split(".")
            parent = None
            for k in range(len(parts) - 1, 0, -1):
                cand = ".".join(parts[:k])
                if cand in by_code:
                    parent = cand
                    break
            item = {"id": code, "code": code, "name": name, "level": len(parts), "parent": parent,
                    "coded": True, "row": r}
            by_code[code] = item
            last_coded = item
        else:
            assert last_coded is not None, f"строка {r} без кода до первого кода"
            pc = last_coded["code"]
            sub_counter[pc] = sub_counter.get(pc, 0) + 1
            item = {"id": f"{pc}-{sub_counter[pc]}", "code": None, "name": name,
                    "level": last_coded["level"] + 1, "parent": pc, "coded": False, "row": r}
        item["object_types"] = applies
        item["expert_note"] = notes.get(r)
        rows.append(item)
    return rows


if __name__ == "__main__":
    src = Path(sys.argv[1])
    rows = parse(src)
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    if out:
        out.write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
    for it in rows:
        ot = "".join("1" if k in it["object_types"] else "." for k, _ in OBJECT_TYPES)
        print(f'{it["row"]:>3} {it["id"]:<14} L{it["level"]} p={it["parent"] or "-":<8} {ot} {it["name"][:90]}' + (f'  # {it["expert_note"][:60]}' if it["expert_note"] else ""))
    print(len(rows), "rows")
